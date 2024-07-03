from __future__ import division
import os
import sys
module_directory = r"C:/Users/covel/OneDrive/Desktop/Rerun-for-Anna/Dom Plotting Code"

# Add the directory to sys.path
if module_directory not in sys.path:
    sys.path.append(module_directory)
import copy
import numpy as np
from numpy import ma
import math as m
import pandas as pd
import matplotlib.patches as mpatches
import string
import cmasher as cmr
import docx

import pickle
from numpy import genfromtxt
from scipy.interpolate import griddata
from matplotlib import rc
import matplotlib.pyplot as plt
plt.style.use('default')
import matplotlib.gridspec as gridspec
import matplotlib.ticker as ticker
import matplotlib.colors as colors
from matplotlib.pyplot import figure
from matplotlib.lines import Line2D

import datetime
from matplotlib.dates import DayLocator, MonthLocator, HourLocator, AutoDateLocator, DateFormatter, drange
from matplotlib.dates import MO, TU, WE, TH, FR, SA, SU, WeekdayLocator
from numpy import arange
from matplotlib.ticker import ScalarFormatter
from matplotlib.ticker import FormatStrFormatter
from matplotlib.ticker import NullFormatter
from mpl_toolkits.axes_grid1 import make_axes_locatable

import matplotlib.cm as cm
import matplotlib.mlab as mlab

import glob

# Font / Figure Size
import matplotlib.pylab as pylab
params = {'legend.fontsize': 'xx-large',
         'axes.labelsize': 'xx-large',
         'axes.titlesize': 'xx-large',
         'xtick.labelsize': 'x-large',
         'ytick.labelsize': 'x-large'}
pylab.rcParams.update(params)
#rc('font',**{'family':'sans-serif','sans-serif':['Arial']})
#rc('text', usetex=True)

pylab.rcParams['xtick.major.pad']='12'
pylab.rcParams['ytick.major.pad']='12'
pylab.rcParams['axes.titlepad'] = 35
pylab.rcParams['axes.labelpad'] = 15

# Custom Colors

darkblue = '#4e79a7'
orange = '#f28e2b'
red = '#e15759'
lightblue = '#76b7b2'
green = '#59a14f'
yellow = '#edc948'
purple = '#b07aa1'
pink = '#ff9da7'
brown = '#9c755f'
gray = '#bab0ac'

# Make Legend
solar = mpatches.Patch(color=yellow, label='Solar')
wind = mpatches.Patch(color='cornflowerblue', label='Wind')
batt = mpatches.Patch(color=purple, label='Li-ion')
rfb = mpatches.Patch(color=green, label='RFB')
caes = mpatches.Patch(color=lightblue, label='CAES')
psh = mpatches.Patch(color=darkblue, label='PSH')
thermal = mpatches.Patch(color=orange, label='Thermal')
gravitational = mpatches.Patch(color=red, label='Gravitational')
metal_air = mpatches.Patch(color=gray, label='Metal-Air')
pgp = mpatches.Patch(color=pink, label='PGP')

from extract_data_sep_power import get_data_sep_power
from extract_data_one_power import get_data_one_power
from extract_data_non_sep_power import get_data_non_sep_power

from math import log10, floor
def round_to_1(x):
    return round(x, -int(floor(log10(abs(x)))))

Li_ion_PGP_X = {}
Li_ion_X = {}
PGP_X = {}
X = {}

Li_ion_PGP_X['directory'] = 'C:/Users/covel/OneDrive/Desktop/Rerun-for-Anna/Output_Data/For Figs 3 and 4/ISNE/three_storage'
Li_ion_X['directory'] = 'C:/Users/covel/OneDrive/Desktop/Rerun-for-Anna/Output_Data/For Figs 3 and 4/ISNE/two_storage_lion'
PGP_X['directory'] = 'C:/Users/covel/OneDrive/Desktop/Rerun-for-Anna/Output_Data/For Figs 3 and 4/ISNE/two_storage_pgp'
X['directory'] = 'C:/Users/covel/OneDrive/Desktop/Rerun-for-Anna/Output_Data/For Figs 3 and 4/ISNE/one_storage'

Li_ion_PGP_X['techs'] = ['RFB', 'PSH', 'Gravitational', 'Thermal', 'CAES', 'Metal-Air']
Li_ion_X['techs'] = ['RFB', 'Gravitational', 'PSH', 'CAES','Thermal', 'Metal-Air', 'PGP']
PGP_X['techs'] = ['RFB', 'Gravitational', 'PSH', 'CAES','Thermal', 'Metal-Air', 'PGP']
X['techs'] = ['Li-ion','RFB',  'Gravitational', 'PSH', 'CAES','Thermal', 'Metal-Air', 'PGP']

def create_dataframe(case, keys):
    path = case['directory']
    case['data'] = {}
    case['disp_data'] = {}
    dataframe = case['disp_data']
    for tech in case['techs']:
        case['data'][tech] = {}
        dic = case['data'][tech]
        if tech == 'Metal-Air':
            
            dic = get_data_non_sep_power(path + '/' + tech, 2, 'fixed_cost', 3, 'fixed_cost', False)
            dic['to_third_tech_cost'] = [None]
            dic['to_third_tech_cap'] = [None]
            dic['to_third_tech_tot_dispatch'] = [None]
            
            dic['third_tech_energy_cost'] = dic.pop('third_tech_cost')
            dic['third_tech_energy_cap'] = dic.pop('third_tech_cap')
            dic['third_tech_energy_tot_dispatch'] = dic.pop('third_tech_tot_dispatch')
            
            dic['from_third_tech_cost'] = [None]
            dic['from_third_tech_cap'] = [None]
            dic['from_third_tech_tot_dispatch'] = [None]
            
            dic['third_tech_dur'] = [100]
            
        elif tech == 'CAES':
            dic = get_data_sep_power(path + '/' + tech, 2, 'fixed_cost', 3, 'fixed_cost', False)
        else:
            dic = get_data_one_power(path + '/' + tech, 2, 'fixed_cost', 3, 'fixed_cost', True)
            dic['to_third_tech_cost'] = dic.pop('third_tech_power_cost')
            dic['to_third_tech_cap'] = dic.pop('third_tech_power_cap')
            dic['to_third_tech_tot_dispatch'] = dic.pop('third_tech_power_tot_dispatch')
            dic['from_third_tech_cost'] = [None]
            dic['from_third_tech_cap'] = [None]
            dic['from_third_tech_tot_dispatch'] = [None]
        data_you_want = {key: None if ((len(dic[key])==0) or (dic[key][0] is None)) else round(dic[key][0], 3) for key in keys}
        dataframe[tech] = pd.DataFrame.from_dict(data_you_want, orient='index', columns=[tech])

keys = ['system_cost', 
        'battery_cap',
        'third_tech_dur', 'to_third_tech_cap', 'third_tech_energy_cap', 'from_third_tech_cap',
        'PGP_dur', 'to_PGP_cap', 'PGP_storage_cap', 'from_PGP_cap']

def concat_dataframe(case):
    df = case['disp_data']
    df1 = df[case['techs'][0]]
    df2 = df[case['techs'][1]]
    df3 = df[case['techs'][2]]
    df4 = df[case['techs'][3]]
    df5 = df[case['techs'][4]]
    df6 = df[case['techs'][5]]
    #df7 = df[case['techs'][6]]
    #df8 = df[case['techs'][7]]
    final_df = pd.concat([df1.set_index(df6.index), df2.set_index(df6.index), 
               df3.set_index(df6.index), df4.set_index(df6.index),
               df5.set_index(df6.index), df6.set_index(df6.index)], axis=1)
    return(final_df)

def save_table(dataframe):
    df = concat_dataframe(dataframe)
    doc = docx.Document('./Tables.docx')

    # add a table to the end and create a reference variable
    # extra row is so we can add the header row
    t = doc.add_table(df.shape[0]+1, df.shape[1])

    # add the header rows.
    for j in range(df.shape[-1]):
        t.cell(0,j).text = df.columns[j]

    # add the rest of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i+1,j).text = str(df.values[i,j])

    # save the doc
    doc.save('./Tables.docx')

keys = ['system_cost', 
        'battery_cap',
        'to_third_tech_cap', 'third_tech_energy_cap', 'from_third_tech_cap',
        'to_PGP_cap', 'PGP_storage_cap', 'from_PGP_cap', 
        'third_tech_dur', 'PGP_dur',
        'batt_cycles', 'third_tech_cycles', 'PGP_cycles']

case = Li_ion_PGP_X

create_dataframe(case, keys)
result = concat_dataframe(case)
display(result)
result.to_csv('C:/Users/covel/OneDrive/Desktop/Rerun-for-Anna/Li_ion_PGP_X_ISNE.csv', index=True)